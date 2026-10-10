"""Zitieren (ragcite): Fundstelle (Abschnitt, Seite, Absatz), Literaturangaben in
Zitierstilen, Vorbelegung beim Indexieren, Anzeige und Einstellungen.

Kein echter Anbieter- oder Crossref-Aufruf: Embeddings gemockt, Crossref über respx.
"""

import io
import json
import re
from unittest import mock

import httpx
import pytest
import respx
from django.urls import reverse

from multigpt.accounts.models import Role, User, UserGroup
from multigpt.chat import citations, services, sources
from multigpt.chat.citations import Locator, Reference, citation_label
from multigpt.chat.forms_citation import DocumentCitationForm
from multigpt.chat.models import (
    EMBEDDING_DIMENSIONS,
    Chunk,
    Collection,
    Conversation,
    Document,
    Message,
    RagSettings,
    Share,
    SourceRef,
)
from multigpt.chat.rag import chat as rag_chat
from multigpt.chat.rag import chunking, crossref, extract, ingest, search
from multigpt.chat.rag.extract import Page

PASSWORD = "Geheim-Test-1234"


# --- Hilfen ------------------------------------------------------------------------


def words(n, prefix="w"):
    return " ".join(f"{prefix}{i:03d}" for i in range(n))


def make_user(username, role="adult"):
    return User.objects.create_user(username, password=PASSWORD, role=Role.objects.get(key=role))


def axis():
    return [1.0] + [0.0] * (EMBEDDING_DIMENSIONS - 1)


def fake_vectors(texts):
    return [axis() for _ in texts]


# --- Absätze und Seiten beim Zerteilen ------------------------------------------------


def test_paragraphs_counted_per_page_and_reset():
    pages = [Page(1, "Eins.\n\nZwei.\n\n\n\nDrei."), Page(2, "Vier.\n\nFünf.")]
    chunks = chunking.split_pages(pages, chunk_tokens=2, overlap_tokens=0)
    found = [(c.text, c.page, c.paragraph) for c in chunks]
    assert found == [
        ("Eins.", 1, 1),
        ("Zwei.", 1, 2),
        ("Drei.", 1, 3),  # mehrere Leerzeilen = eine Grenze
        ("Vier.", 2, 1),  # neue Seite: wieder ab 1
        ("Fünf.", 2, 2),
    ]


def test_paragraphs_continue_without_pages():
    pages = [Page(None, "A a.\n\nB b."), Page(None, "C c.")]
    chunks = chunking.split_pages(pages, chunk_tokens=2, overlap_tokens=0)
    assert [c.paragraph for c in chunks] == [1, 2, 3]
    assert all(c.page is None and c.page_end is None for c in chunks)


def test_page_change_ends_paragraph_and_ranges():
    # Ein Abschnitt über Seitengrenze: erste Seite, letzte Seite, Absatz von–bis.
    pages = [Page(12, "x\n\nAbsatz drei geht bis zum Seitenende"), Page(13, "Weiter auf Seite.")]
    [chunk] = chunking.split_pages(pages, chunk_tokens=500, overlap_tokens=0)
    assert (chunk.page, chunk.page_end) == (12, 13)
    assert (chunk.paragraph, chunk.paragraph_end) == (1, 1)
    assert "Seitenende\n\nWeiter" in chunk.text  # Seitenwechsel = Absatzgrenze


def test_page_is_first_page_not_majority():
    pages = [Page(1, "kurz"), Page(2, words(40))]
    [chunk] = chunking.split_pages(pages, chunk_tokens=500, overlap_tokens=0)
    assert (chunk.page, chunk.page_end) == (1, 2)


def test_overlap_starts_in_paragraph_of_first_segment():
    text = words(30, "a") + "\n\n" + words(30, "b")
    chunks = chunking.split_pages([Page(5, text)], chunk_tokens=40, overlap_tokens=10)
    assert len(chunks) >= 2
    second = chunks[1]
    # Die Überlappung beginnt mitten in Absatz 1; es zählt der Absatz des ersten Wortes.
    assert second.text.startswith("a")
    assert (second.paragraph, second.paragraph_end) == (1, 2)
    assert chunks[-1].paragraph_end == 2


def test_heading_pages_and_section_but_no_paragraph():
    assert chunking.heading("Bericht", page=12, page_end=13) == "Dokument: Bericht\nSeiten: 12–13"
    assert chunking.heading("Bericht", page=12, page_end=12) == "Dokument: Bericht\nSeite: 12"
    text = chunking.heading("DIN EN ISO 9001", page=3, section="7.5.3", section_title="Lenkung")
    assert text.splitlines()[-1] == "Abschnitt: 7.5.3 Lenkung"
    assert "Abs." not in text


# --- Gliederung (Abschnittsnummern) ----------------------------------------------------

NORM_TOC = (
    "Inhalt\n\n0 Einleitung 4\n1 Anwendungsbereich 5\n2 Normative Verweisungen 5\n"
    "7 Unterstützung 8\n7.5 Dokumentierte Information ........ 9\nAnhang A (informativ) 20"
)
NORM_BODY = (
    "Vorwort\n\nDies ist ein Vorwort.\n\n"
    "0 Einleitung\n\nText der Einleitung.\n\n"
    "1 Anwendungsbereich\n\nDiese Norm gilt für alles.\n"
    "2 Mitarbeiter sind zuständig\nund weiter.\n\n"
    "2 Normative Verweisungen\n\nKeine.\n\n"
    "3 Begriffe\n\nBegriffe stehen hier.\n\n"
    "1. Aufzählung eins\n\n2 Stück\n\n"
    "3.1 Organisation\n\nText.\n\n"
    "4 Kontext\n\n"
    "4.1 Verstehen der Organisation\nDer Text geht direkt weiter. Und so weiter.\n"
    "4.1.1 Allgemeines\nInhalt zu 4.1.1.\n\n"
    "12 Monate Laufzeit\n\n"
    "4.2 Erfordernisse\n\nDie Organisation muss bestimmen.\n\n"
    "Anhang A (informativ)\n\nA.1 Allgemeines\n\nAnhangtext."
)


def section_of(chunks, needle):
    """Abschnitt des ersten Abschnitts, der mit ``needle`` beginnt."""
    return next(c.section for c in chunks if c.text.startswith(needle))


def test_sections_nested_annex_and_toc_skipped():
    pages = [Page(1, NORM_TOC), Page(2, NORM_BODY)]
    # Ein Wort je Abschnitt: so ist die Gliederung jedes Wortes sichtbar.
    chunks = chunking.split_pages(pages, chunk_tokens=1, overlap_tokens=0)
    # Inhaltsverzeichnis: keine Abschnitte.
    assert all(c.section == "" for c in chunks if c.page == 1)
    assert section_of(chunks, "Vorwort.") == ""
    assert section_of(chunks, "Einleitung.") == "0"
    assert section_of(chunks, "alles.") == "1"
    assert section_of(chunks, "Keine.") == "2"  # „2 Mitarbeiter … und weiter“ ist keine Überschrift
    assert section_of(chunks, "stehen") == "3"
    assert section_of(chunks, "Stück") == "3"  # „1. Aufzählung“ und „2 Stück“ sind keine
    assert section_of(chunks, "Text.") == "3.1"
    assert section_of(chunks, "direkt") == "4.1"  # Überschrift ohne Leerzeile
    assert section_of(chunks, "4.1.1.") == "4.1.1"
    assert section_of(chunks, "Laufzeit") == "4.1.1"  # „12 Monate“ passt nicht
    assert section_of(chunks, "bestimmen.") == "4.2"
    assert section_of(chunks, "Anhangtext") == "A.1"
    titled = next(c for c in chunks if c.text.startswith("bestimmen."))
    assert titled.section_title == "Erfordernisse"


def test_section_end_only_when_crossing():
    text = "1 Einleitung\n\nText eins.\n\n2 Zweck\n\nText zwei."
    [chunk] = chunking.split_pages([Page(1, text)], chunk_tokens=500, overlap_tokens=0)
    assert (chunk.section, chunk.section_end) == ("1", "2")
    chunks = chunking.split_pages([Page(1, text)], chunk_tokens=4, overlap_tokens=0)
    assert all(c.section_end == "" for c in chunks if c.section == "2")


def test_sections_resync_after_missed_heading():
    # „5 …“ fehlt (nicht erkannt); 6 und 6.1 folgen aufeinander -> wieder aufgesetzt.
    text = "1 Eins\n\nA.\n\n2 Zwei\n\nB.\n\n6 Sechs\n\nC.\n\n6.1 Teil\n\nD."
    chunks = chunking.split_pages([Page(1, text)], chunk_tokens=1, overlap_tokens=0)
    assert section_of(chunks, "C.") == "6"
    assert section_of(chunks, "D.") == "6.1"


# --- Absätze aus den Formaten ----------------------------------------------------------


def test_docx_one_paragraph_per_w_p(tmp_path):
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("Überschrift", level=1)
    doc.add_paragraph("Erster Absatz.")
    doc.add_paragraph("")  # leer: zählt nicht
    doc.add_paragraph("Zweiter Absatz.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "A", "B"
    table.cell(1, 0).text, table.cell(1, 1).text = "C", "D"
    doc.add_paragraph("Nach der Tabelle.")
    path = tmp_path / "t.docx"
    doc.save(path)
    [page] = extract.extract(str(path), extract.KIND_DOCX)
    assert page.text == (
        "Überschrift\n\nErster Absatz.\n\nZweiter Absatz.\n\nA | B\nC | D\n\nNach der Tabelle."
    )
    chunks = chunking.split_pages([page], chunk_tokens=3, overlap_tokens=0)
    assert next(c.paragraph for c in chunks if "Nach" in c.text) == 5


PDF_TEXT = (
    "Einleitung\n"
    "Dies ist der erste Absatz mit einigen Sätzen. Er geht über mehrere Zeilen, damit man\n"
    "sieht, wie der Umbruch aussieht und ob die Textebene Leerzeilen liefert. Noch ein\n"
    "Satz hier.\n"
    "Zweiter Absatz beginnt hier. Auch er ist länger und hat mehrere Zeilen Text, damit\n"
    "der Umbruch innerhalb des Absatzes sichtbar wird. Ende des zweiten Absatzes ist\n"
    "erreicht.\n"
    "Dritter Absatz kurz."
)


def test_pdf_paragraph_heuristic():
    result = extract.pdf_paragraphs(PDF_TEXT)
    blocks = result.split("\n\n")
    assert [b.split("\n")[0][:12] for b in blocks] == [
        "Einleitung",
        "Dies ist der",
        "Zweiter Absa",
        "Dritter Absa",
    ]
    # Volle Zeilen mit Satzende im Absatz bleiben zusammen.
    assert "Noch ein\nSatz hier." in result


def test_pdf_paragraph_heuristic_leaves_existing_and_short_text():
    assert extract.pdf_paragraphs("A.\n\nB.") == "A.\n\nB."
    assert extract.pdf_paragraphs("Kurz.\nNoch kurz.") == "Kurz.\nNoch kurz."


def test_olmocr_table_is_own_paragraph():
    from multigpt.chat.rag import ocr

    page = ocr.parse_olmocr("Text davor.\n<table><tr><td>A</td><td>B</td></tr></table>\nDanach.")
    assert "Text davor.\n\nA | B\n\nDanach." in extract.clean_text(page.text)


# --- citation_label ---------------------------------------------------------------------


@pytest.mark.parametrize(
    ("args", "kwargs", "expected"),
    [
        ((12, 12, 3, 3), {}, "S. 12, Abs. 3"),
        ((12, None, 3, None), {}, "S. 12, Abs. 3"),
        ((12, 12, 3, 5), {}, "S. 12, Abs. 3–5"),
        ((12, 13, 4, 2), {}, "S. 12, Abs. 4 – S. 13, Abs. 2"),
        ((12, 13, None, None), {}, "S. 12–13"),
        ((12, None, None, None), {}, "S. 12"),
        ((None, None, 17, 17), {}, "Abs. 17"),
        ((None, None, 17, 19), {}, "Abs. 17–19"),
        ((None, None, None, None), {}, ""),
        ((12, 12, 3, 3), {"section": "7.5.3"}, "Abschn. 7.5.3, S. 12, Abs. 3"),
        ((None, None, None, None), {"section": "7.5", "section_end": "7.6"}, "Abschn. 7.5–7.6"),
        ((12, 12, 3, 3), {"page_prefix": ""}, "12, Abs. 3"),
    ],
)
def test_citation_label(args, kwargs, expected):
    assert citation_label(*args, **kwargs) == expected


# --- Zitierstile ------------------------------------------------------------------------

BOOK = Reference(
    title="Grundlagen der Informatik",
    type="book",
    authors=("Müller, Hans Peter", "Schmidt, Eva"),
    date="2024",
    publisher="Springer",
    place="Berlin",
    edition="3",
    isbn="978-3-16-148410-0",
)
ARTICLE = Reference(
    title="Absätze zählen",
    type="article",
    authors=("Müller, Hans", "Schmidt, Eva", "Weber, Karl"),
    date="2023",
    journal="Zeitschrift für Text",
    volume="12",
    issue="3",
    pages="45-67",
    doi="10.1000/xyz123",
)
CHAPTER = Reference(
    title="Kapitel eins",
    type="chapter",
    authors=("Weber, Karl",),
    editors=("Müller, Hans", "Schmidt, Eva"),
    container="Handbuch Textanalyse",
    date="2021",
    publisher="Springer",
    place="Berlin",
    pages="45–67",
    series="Lecture Notes in Computer Science",
    series_number="1234",
    doi="10.1007/978-3-030-12345-6_3",
)
EDITED = Reference(
    title="Handbuch Textanalyse", type="edited", editors=("Müller, Hans",), date="2021"
)
CONFERENCE = Reference(
    title="Fast Search",
    type="conference",
    authors=("Lee, Ann",),
    container="Proceedings of SIGIR 2020",
    publisher="ACM",
    pages="1-10",
    date="2020",
)
NORM = Reference(
    title="Qualitätsmanagementsysteme – Anforderungen",
    type="standard",
    number="DIN EN ISO 9001",
    date="2015-11",
    institution="DIN",
    publisher="DIN Media",
    place="Berlin",
)
WEB = Reference(
    title="Wetter heute", type="web", url="https://example.org/w", accessed="2026-10-10"
)
BARE = Reference(title="Jahresbericht 2024")
LOC = Locator(12, 12, 3, None)

EXPECTED_ENTRY = {
    ("din", "book"): "MÜLLER, Hans Peter; SCHMIDT, Eva, 2024. Grundlagen der Informatik. 3. Aufl. "
    "Berlin: Springer. ISBN 978-3-16-148410-0.",
    ("apa", "book"): "Müller, H. P., & Schmidt, E. (2024). Grundlagen der Informatik (3. Aufl.). "
    "Springer.",
    ("harvard", "book"): "Müller, H. P. und Schmidt, E. (2024) Grundlagen der Informatik. "
    "3. Aufl. Berlin: Springer.",
    ("chicago", "book"): "Müller, Hans Peter, und Eva Schmidt. 2024. Grundlagen der Informatik. "
    "3. Aufl. Berlin: Springer.",
    ("mla", "book"): "Müller, Hans Peter, und Eva Schmidt. Grundlagen der Informatik. 3. Aufl., "
    "Springer, 2024.",
    ("din", "article"): "MÜLLER, Hans; SCHMIDT, Eva; WEBER, Karl, 2023. Absätze zählen. "
    "Zeitschrift für Text. 12(3), S. 45–67. DOI: 10.1000/xyz123.",
    ("apa", "article"): "Müller, H., Schmidt, E., & Weber, K. (2023). Absätze zählen. "
    "Zeitschrift für Text, 12(3), 45–67. https://doi.org/10.1000/xyz123",
    ("din", "chapter"): "WEBER, Karl, 2021. Kapitel eins. In: MÜLLER, Hans; SCHMIDT, Eva (Hrsg.): "
    "Handbuch Textanalyse. Lecture Notes in Computer Science, Bd. 1234. Berlin: Springer, "
    "S. 45–67. DOI: 10.1007/978-3-030-12345-6_3.",
    ("apa", "chapter"): "Weber, K. (2021). Kapitel eins. In H. Müller & E. Schmidt (Hrsg.), "
    "Handbuch Textanalyse (S. 45–67). Springer. https://doi.org/10.1007/978-3-030-12345-6_3",
    ("din", "standard"): "DIN EN ISO 9001:2015-11. Qualitätsmanagementsysteme – Anforderungen. "
    "Berlin: DIN Media.",
    ("apa", "standard"): "DIN. (2015). Qualitätsmanagementsysteme – Anforderungen "
    "(DIN EN ISO 9001:2015-11). DIN Media.",
    ("din", "web"): "Wetter heute [online], o. J. [Zugriff am: 10.10.2026]. "
    "Verfügbar unter: https://example.org/w",
    (
        "apa",
        "web",
    ): "Wetter heute. (o. J.). Abgerufen am 10. Oktober 2026, von https://example.org/w",
    ("mla", "web"): "„Wetter heute.“ https://example.org/w. Abgerufen am 10. Okt. 2026.",
}
REFS = {
    "book": BOOK,
    "article": ARTICLE,
    "chapter": CHAPTER,
    "standard": NORM,
    "web": WEB,
}


@pytest.mark.parametrize(("key", "expected"), EXPECTED_ENTRY.items())
def test_entries(key, expected):
    style, kind = key
    assert citations.entry(REFS[kind], style) == expected


@pytest.mark.parametrize(
    ("ref", "style", "expected"),
    [
        (BOOK, "din", "(Müller und Schmidt 2024, S. 12, Abs. 3)"),
        (BOOK, "apa", "(Müller & Schmidt, 2024, S. 12, Abs. 3)"),
        (BOOK, "harvard", "(Müller und Schmidt 2024, S. 12, Abs. 3)"),
        (BOOK, "chicago", "(Müller und Schmidt 2024, 12, Abs. 3)"),
        (BOOK, "mla", "(Müller und Schmidt 12, Abs. 3)"),
        (ARTICLE, "din", "(Müller u. a. 2023, S. 12, Abs. 3)"),
        (ARTICLE, "apa", "(Müller et al., 2023, S. 12, Abs. 3)"),
        (NORM, "apa", "(DIN EN ISO 9001:2015-11, S. 12, Abs. 3)"),
        (NORM, "mla", "(DIN EN ISO 9001:2015-11, S. 12, Abs. 3)"),
        (BARE, "din", "(Jahresbericht 2024 o. J., S. 12, Abs. 3)"),
        (BARE, "mla", "(Jahresbericht 2024 12, Abs. 3)"),
    ],
)
def test_short_citations(ref, style, expected):
    assert citations.short(ref, style, LOC) == expected


def test_short_with_section_for_norm():
    where = Locator(12, 12, None, None, "7.5.3")
    assert citations.short(NORM, "din", where) == "(DIN EN ISO 9001:2015-11, Abschn. 7.5.3, S. 12)"
    assert citations.short(BOOK, "apa") == "(Müller & Schmidt, 2024)"


@pytest.mark.parametrize("style", citations.STYLES)
@pytest.mark.parametrize("ref", [BOOK, ARTICLE, CHAPTER, EDITED, CONFERENCE, NORM, WEB, BARE])
def test_every_style_and_type_without_placeholders(style, ref):
    text = citations.entry(ref, style)
    assert text and ref.title in text
    assert "o. A." not in text and "None" not in text
    assert "  " not in text and ".." not in text.replace("…", "")
    assert not re.search(r"\(\s*\)|,\s*[,.]", text)
    assert citations.short(ref, style).startswith("(")


@pytest.mark.parametrize("style", citations.STYLES)
def test_incomplete_data_is_left_out(style):
    ref = Reference(title="Nur Titel", type="chapter", container="Sammelwerk")
    text = citations.entry(ref, style)
    assert "Hrsg." not in text and "S." not in text.replace("Sammelwerk", "")
    if style == "mla":
        assert "o. J." not in text  # MLA lässt fehlendes Jahr weg
    else:
        assert "o. J." in text


def test_edited_volume_puts_editors_first():
    assert citations.entry(EDITED, "din").startswith("MÜLLER, Hans (Hrsg.), 2021.")
    assert citations.entry(EDITED, "apa").startswith("Müller, H. (Hrsg.). (2021).")
    assert citations.entry(EDITED, "mla").startswith("Müller, Hans, Hrsg.")


def test_norm_with_status_and_replacement():
    ref = Reference.from_dict(
        {**NORM.to_dict(), "status": "withdrawn", "replaces": "DIN EN ISO 9001:2008-12"}
    )
    text = citations.entry(ref, "din")
    assert text.endswith("Zurückgezogen. Ersatz für DIN EN ISO 9001:2008-12.")


def test_bibtex_types():
    chapter = citations.bibtex(CHAPTER)
    assert chapter.startswith("@incollection{weber2021kapitel,")
    assert "editor = {Müller, Hans and Schmidt, Eva}" in chapter
    assert "booktitle = {Handbuch Textanalyse}" in chapter
    assert "pages = {45--67}" in chapter and "volume = {1234}" in chapter
    norm = citations.bibtex(NORM)
    assert norm.startswith("@techreport{dineniso90012015,")
    assert "number = {DIN EN ISO 9001:2015-11}" in norm and "type = {Norm}" in norm
    assert citations.bibtex(CONFERENCE).startswith("@inproceedings{")
    web = citations.bibtex(WEB)
    assert web.startswith("@online{") and "urldate = {2026-10-10}" in web
    corp = citations.bibtex(Reference(title="A & B 100%", authors=("Statistisches Bundesamt",)))
    assert "author = {{Statistisches Bundesamt}}" in corp and r"A \& B 100\%" in corp


def test_reference_roundtrip_and_robust_from_dict():
    assert Reference.from_dict(CHAPTER.to_dict()) == CHAPTER
    odd = Reference.from_dict({"type": "quatsch", "authors": "kein Array", "title": 5})
    assert (odd.type, odd.authors, odd.title) == ("other", (), "5")


# --- Erkennung -------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("DIN_EN_ISO_9001_2015-11.pdf", ("DIN EN ISO 9001", "2015-11", "DIN")),
        ("DIN EN ISO 9001:2015-11 Qualitätsmanagement", ("DIN EN ISO 9001", "2015-11", "DIN")),
        ("DIN 1450 Ausgabe 2013-04", ("DIN 1450", "2013-04", "DIN")),
        ("Norm DIN 1450, Ausgabe: 2013-04", ("DIN 1450", "2013-04", "DIN")),
        ("ISO/IEC 27001:2022-10", ("ISO/IEC 27001", "2022-10", "ISO")),
        ("VDI 2230-1:2015-11", ("VDI 2230-1", "2015-11", "VDI")),
        ("DIN EN ISO 9001 ohne Datum", None),
        ("Tabelle 2015-11", None),
        ("DIN 1450:2013-13", None),
    ],
)
def test_find_norm(text, expected):
    found = citations.find_norm(text)
    result = (found.number, found.date, found.institution) if found else None
    assert result == expected


def test_find_doi():
    assert citations.find_doi("DOI: https://doi.org/10.1007/978-3-030-12345-6_3.") == (
        "10.1007/978-3-030-12345-6_3"
    )
    assert citations.find_doi("ohne") == ""


# --- Validierung --------------------------------------------------------------------------


def form(**data):
    return DocumentCitationForm(data={"bib_type": "book", **data}, instance=Document())


def test_form_normalizes_and_validates():
    ok = form(
        bib_date="12.03.2024", bib_doi="https://doi.org/10.1000/ABC", bib_isbn="3-16-148410-X"
    )
    ok.is_valid()
    assert ok.cleaned_data["bib_date"] == "2024-03-12"
    assert ok.cleaned_data["bib_doi"] == "10.1000/ABC"
    assert "bib_isbn" not in ok.errors and "bib_date" not in ok.errors
    bad = form(
        bib_date="2024-13",
        bib_isbn="978-3-16-148410-1",
        bib_doi="11.1/x",
        bib_url="javascript:alert(1)",
        bib_pages="45 bis 67",
    )
    assert not bad.is_valid()
    assert set(bad.errors) >= {"bib_date", "bib_isbn", "bib_doi", "bib_url", "bib_pages"}


# --- Kontextblock und Werkzeug -------------------------------------------------------------


def test_context_block_has_citable_location_and_short():
    entry = sources.ContextEntry(
        n=2,
        kind=SourceRef.Kind.DOCUMENT,
        title="Jahresbericht 2024",
        text="Umsatz stieg.",
        page=12,
        page_end=12,
        paragraph=3,
        paragraph_end=5,
        short="(Müller 2024, S. 12, Abs. 3–5)",
    )
    block = sources.context_block([entry])
    assert 'seite="12" fundstelle="S. 12, Abs. 3–5" kurzbeleg="(Müller 2024, S. 12, Abs. 3–5)"' in (
        block
    )
    assert "\n[2] Jahresbericht 2024, S. 12, Abs. 3–5\n\nUmsatz stieg.\n" in block
    assert "Seite und Absatz" in sources.SYSTEM_NOTE


@pytest.fixture
def anna(db):
    user = make_user("anna")
    user.citation_style = "apa"
    user.save()
    return user


@pytest.fixture
def doc(anna):
    coll = Collection.objects.create(owner=anna, name="Normen")
    return Document.objects.create(
        collection=coll,
        title="ISO 9001",
        file="documents/x.pdf",
        status=Document.Status.INDEXED,
        bib_type="standard",
        bib_number="DIN EN ISO 9001",
        bib_date="2015-11",
        bib_institution="DIN",
        bib_title="Qualitätsmanagementsysteme",
    )


@pytest.fixture
def norm_chunk(doc):
    return Chunk.objects.create(
        document=doc,
        position=0,
        text="Die Organisation muss dokumentierte Information lenken.",
        page=12,
        page_end=12,
        paragraph=3,
        paragraph_end=4,
        section="7.5.3",
        section_title="Lenkung",
        embedding=axis(),
    )


def test_search_documents_tool_output(anna, norm_chunk, monkeypatch, settings):
    settings.DEBUG = True
    settings.RAG_FAKE_EMBEDDINGS = True
    monkeypatch.setattr(search, "embed_query", lambda text: axis())
    conv = Conversation.objects.create(user=anna)
    answer = services.append_message(conv, role=Message.Role.ASSISTANT, content="")
    collector = sources.SourceCollector(answer)
    result = rag_chat.run_search_documents(anna, {"query": "lenken"}, collector)
    assert 'fundstelle="Abschn. 7.5.3, S. 12, Abs. 3–4"' in result.text
    assert 'kurzbeleg="(DIN EN ISO 9001:2015-11, Abschn. 7.5.3, S. 12, Abs. 3–4)"' in result.text
    ref = SourceRef.objects.get(message=answer)
    assert (ref.page, ref.page_end, ref.paragraph, ref.paragraph_end, ref.section) == (
        12,
        12,
        3,
        4,
        "7.5.3",
    )
    assert ref.biblio["number"] == "DIN EN ISO 9001"
    # Die Angabe bleibt nach dem Löschen des Dokuments.
    norm_chunk.document.delete()
    ref.refresh_from_db()
    item = sources.serialize(ref, 1, citations.prefs_for(anna))
    assert item["location"] == "Abschn. 7.5.3, S. 12, Abs. 3–4"
    assert item["entry"].startswith("DIN. (2015). Qualitätsmanagementsysteme")


# --- Anzeige: API und Template ------------------------------------------------------------


@pytest.fixture
def answered(anna, norm_chunk):
    conv = Conversation.objects.create(user=anna, title="Normen")
    services.append_message(conv, role=Message.Role.USER, content="Was steht in 7.5.3?")
    answer = services.append_message(conv, role=Message.Role.ASSISTANT, content="Siehe [1].")
    sources.SourceCollector(answer).add(
        SourceRef.Kind.DOCUMENT,
        norm_chunk.document.title,
        chunk=norm_chunk,
        page=12,
        page_end=12,
        paragraph=3,
        paragraph_end=4,
        section="7.5.3",
        biblio=citations.reference_from_document(norm_chunk.document).to_dict(),
    )
    SourceRef.objects.create(
        message=answer, kind=SourceRef.Kind.WEB, title="Wetter <b>", url="https://ex.org/w"
    )
    return conv


def test_api_messages_return_citations_in_viewer_style(client, anna, answered):
    client.force_login(anna)
    data = client.get(reverse("chat:api_messages", args=[answered.pk])).json()
    doc_source, web_source = data[-1]["sources"]
    assert doc_source["style"] == "apa"
    assert doc_source["location"] == "Abschn. 7.5.3, S. 12, Abs. 3–4"
    assert doc_source["short"] == "(DIN EN ISO 9001:2015-11, Abschn. 7.5.3, S. 12, Abs. 3–4)"
    assert doc_source["entry"].startswith("DIN. (2015). Qualitätsmanagementsysteme")
    assert set(doc_source["formats"]) == {"din", "apa", "harvard", "chicago", "mla", "bibtex"}
    assert doc_source["formats"]["din"]["entry"].startswith("DIN EN ISO 9001:2015-11.")
    assert web_source["formats"]["bibtex"].startswith("@online{")
    assert "Abgerufen am" in web_source["entry"]  # Abrufdatum = Zeitpunkt der Antwort


def test_locator_switch_hides_location(client, anna, answered):
    anna.citation_locator = False
    anna.save()
    client.force_login(anna)
    doc_source = client.get(reverse("chat:api_messages", args=[answered.pk])).json()[-1]["sources"][
        0
    ]
    assert doc_source["location"] == ""
    assert doc_source["short"] == "(DIN EN ISO 9001:2015-11)"


def test_template_renders_styled_entry_and_escaped_data(client, anna, answered):
    client.force_login(anna)
    html = client.get(reverse("chat:conversation", args=[answered.pk])).content.decode()
    block = re.search(r'<section class="message-sources".*?</section>', html, re.S).group(0)
    assert "DIN. (2015). Qualitätsmanagementsysteme" in block
    assert '<span class="message-source-location">Abschn. 7.5.3, S. 12, Abs. 3–4</span>' in block
    assert "Wetter &lt;b&gt;" in block and "<b>" not in block
    raw = re.search(r'data-citations="([^"]*)"', block).group(1)
    from html import unescape

    data = json.loads(unescape(raw))
    assert data[0]["short"].startswith("(DIN EN ISO 9001:2015-11")
    assert "chat/citations.js" in html


# --- Einstellungen --------------------------------------------------------------------------


def test_settings_page(client, anna):
    assert client.get(reverse("settings")).status_code == 302  # nicht angemeldet
    client.force_login(anna)
    page = client.get(reverse("settings")).content.decode()
    assert "Zitierstil" in page and "DIN ISO 690" in page
    response = client.post(
        reverse("settings"),
        {"citation_style": "mla", "citation_short": "on", "citation_locator": ""},
    )
    assert response.status_code == 302
    anna.refresh_from_db()
    assert (anna.citation_style, anna.citation_short, anna.citation_locator) == ("mla", True, False)
    assert client.post(reverse("settings"), {"citation_style": "x"}).status_code == 200
    anna.refresh_from_db()
    assert anna.citation_style == "mla"
    assert reverse("settings") in client.get(reverse("chat:index")).content.decode()


def test_default_style_is_din(db):
    user = make_user("neu")
    assert citations.prefs_for(user) == citations.Prefs("din", False, True)


# --- Literaturangaben bearbeiten: Rechte ----------------------------------------------------


def test_citation_edit_rights(client, anna, doc):
    url = reverse("chat:document_citation", args=[doc.pk])
    bernd = make_user("bernd")
    client.force_login(bernd)
    assert client.get(url).status_code == 404  # fremd
    group = UserGroup.objects.create(name="Leser")
    bernd.groups.add(group)
    Share.objects.create(collection=doc.collection, group=group, can_write=False)
    assert client.get(url).status_code == 403  # nur Leser
    assert client.post(url, {"bib_type": "book"}).status_code == 403
    client.force_login(anna)
    page = client.get(url).content.decode()
    assert "DIN EN ISO 9001:2015-11. Qualitätsmanagementsysteme." in page
    response = client.post(
        url,
        {
            "bib_type": "standard",
            "bib_title": "<script>alert(1)</script>",
            "bib_number": "DIN EN ISO 9001",
            "bib_date": "2015-11",
        },
    )
    assert response.status_code == 302
    doc.refresh_from_db()
    assert doc.bib_edited
    page = client.get(url).content.decode()
    assert "<script>alert(1)</script>" not in page and "&lt;script&gt;" in page
    bad = client.post(url, {"bib_type": "book", "bib_isbn": "123"})
    assert bad.status_code == 200 and "Keine gültige ISBN" in bad.content.decode()
    detail = client.get(reverse("chat:collection_detail", args=[doc.collection_id]))
    assert url in detail.content.decode()


# --- Vorbelegung beim Indexieren -------------------------------------------------------------


@pytest.fixture
def media(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path / "media"
    return settings.MEDIA_ROOT


def docx_bytes(title="Ein Bericht über Normen", author="Müller, Hans", text=None):
    from docx import Document as DocxDocument

    document = DocxDocument()
    document.core_properties.title = title
    document.core_properties.author = author
    for paragraph in text or ["1 Einleitung", "Text eins.", "2 Zweck", "Text zwei."]:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload(anna, name, data, title=None):
    coll = Collection.objects.create(owner=anna, name=f"Sammlung {name}")
    doc = Document(collection=coll, title=title or name)
    doc.file.save(name, io.BytesIO(data), save=False)
    doc.save()
    return doc


def test_ingest_stores_location_and_prefills(anna, media):
    doc = upload(anna, "bericht.docx", docx_bytes(), title="DIN_1450_2013-04")
    with mock.patch.object(ingest, "_embed", side_effect=fake_vectors):
        ingest.index_document(doc)
    doc.refresh_from_db()
    assert (doc.bib_type, doc.bib_number, doc.bib_date, doc.bib_institution) == (
        "standard",
        "DIN 1450",
        "2013-04",
        "DIN",
    )
    assert doc.bib_title == "Ein Bericht über Normen"
    assert doc.bib_authors == ""  # Norm: Ersteller der Datei ist kein Autor
    chunk = Chunk.objects.get(document=doc)
    assert (chunk.page, chunk.paragraph, chunk.paragraph_end) == (None, 1, 4)
    assert (chunk.section, chunk.section_end) == ("1", "2")
    assert "Abschnitt: 1 Einleitung" in chunk.heading


def test_ingest_never_overwrites_edited(anna, media):
    doc = upload(anna, "b.docx", docx_bytes())
    doc.bib_edited = True
    doc.save()
    with mock.patch.object(ingest, "_embed", side_effect=fake_vectors):
        ingest.index_document(doc)
    doc.refresh_from_db()
    assert (doc.bib_title, doc.bib_authors) == ("", "")


def test_metadata_junk_title_ignored(anna, media):
    doc = upload(
        anna, "c.docx", docx_bytes(title="Microsoft Word - c.docx", author="Administrator")
    )
    with mock.patch.object(ingest, "_embed", side_effect=fake_vectors):
        ingest.index_document(doc)
    doc.refresh_from_db()
    assert (doc.bib_title, doc.bib_authors, doc.bib_type) == ("", "", "other")


CROSSREF = {
    "message": {
        "type": "book-chapter",
        "title": ["Kapitel aus Crossref"],
        "author": [{"family": "Weber", "given": "Karl"}],
        "editor": [{"family": "Müller", "given": "Hans"}],
        "container-title": ["Handbuch Textanalyse", "Lecture Notes in Computer Science"],
        "publisher": "Springer",
        "publisher-location": "Cham",
        "page": "45-67",
        "issued": {"date-parts": [[2021, 5]]},
        "isbn-type": [
            {"type": "print", "value": "978-3-16-148410-0"},
            {"type": "electronic", "value": "978-3-030-12345-6"},
        ],
    }
}
DOI_TEXT = ["Kapitel", "https://doi.org/10.1007/978-3-030-12345-6_3", "Text."]


def crossref_on(mailto=""):
    cfg = RagSettings.load()
    cfg.crossref_enabled = True
    cfg.crossref_mailto = mailto
    cfg.save()


@respx.mock
def test_crossref_off_makes_no_request(anna, media):
    route = respx.get(url__startswith="https://api.crossref.org/")
    doc = upload(anna, "d.docx", docx_bytes(title="", author="", text=DOI_TEXT))
    with mock.patch.object(ingest, "_embed", side_effect=fake_vectors):
        ingest.index_document(doc)
    assert not route.called
    doc.refresh_from_db()
    assert doc.bib_doi == "10.1007/978-3-030-12345-6_3"  # DOI aus dem Text erkannt


@respx.mock
def test_crossref_on_fills_only_empty_fields(anna, media):
    crossref_on("familie@example.org")
    route = respx.get("https://api.crossref.org/works/10.1007%2F978-3-030-12345-6_3").mock(
        return_value=httpx.Response(200, json=CROSSREF)
    )
    doc = upload(anna, "e.docx", docx_bytes(title="Mein Titel", author="", text=DOI_TEXT))
    doc.bib_publisher = "Eigener Verlag"
    doc.save()
    with mock.patch.object(ingest, "_embed", side_effect=fake_vectors):
        ingest.index_document(doc)
    assert route.called
    assert "mailto:familie@example.org" in route.calls[0].request.headers["User-Agent"]
    doc.refresh_from_db()
    assert doc.bib_type == "chapter"
    assert doc.bib_title == "Kapitel aus Crossref"  # Crossref vor Datei-Metadaten
    assert doc.bib_publisher == "Eigener Verlag"  # vorhandener Wert bleibt
    assert (doc.bib_authors, doc.bib_editors) == ("Weber, Karl", "Müller, Hans")
    assert (doc.bib_container, doc.bib_series) == (
        "Handbuch Textanalyse",
        "Lecture Notes in Computer Science",
    )
    assert (doc.bib_isbn, doc.bib_isbn_e, doc.bib_date) == (
        "978-3-16-148410-0",
        "978-3-030-12345-6",
        "2021-05",
    )


@respx.mock
def test_crossref_errors_are_ignored(anna, media):
    crossref_on()
    respx.get(url__startswith="https://api.crossref.org/").mock(
        side_effect=httpx.ConnectTimeout("zu langsam")
    )
    doc = upload(anna, "f.docx", docx_bytes(title="", author="", text=DOI_TEXT))
    with mock.patch.object(ingest, "_embed", side_effect=fake_vectors):
        ingest.index_document(doc)
    doc.refresh_from_db()
    assert doc.status == Document.Status.INDEXED and doc.bib_type == "other"
    respx.get(url__startswith="https://api.crossref.org/").mock(
        return_value=httpx.Response(200, text="kein json")
    )
    assert crossref.lookup("10.1007/x") is None
    assert crossref.lookup("keine doi") is None
