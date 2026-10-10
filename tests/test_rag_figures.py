"""Abbildungen im RAG (rag.figures): Herauslösen aus PDF und DOCX, Aussortieren,
Grenzen, Bildunterschrift, Beschreibung über ein Vision-Modell, Bilddateien als
Dokument.

Testbilder, PDFs und DOCX entstehen im Test (Pillow, Roh-PDF, python-docx).
Kein echter Anbieteraufruf: das Vision-Modell über respx, Embeddings und OCR
gemockt.
"""

import base64
import io
import json
import zlib
from pathlib import Path
from unittest import mock

import httpx
import pytest
import respx
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from PIL import Image, ImageDraw

from multigpt.accounts.models import Role, User
from multigpt.chat.models import (
    EMBEDDING_DIMENSIONS,
    AIModel,
    Collection,
    Document,
    Job,
    Provider,
    RagSettings,
)
from multigpt.chat.rag import chunking, extract, figures, ingest, jobs
from multigpt.rag import settings_form

pytestmark = pytest.mark.django_db

PASSWORD = "Geheim-Test-1234"
LMSTUDIO = "http://lmstudio.example.invalid:1234/v1"
VISION = "qwen/qwen3-vl-8b"
INTRO = "Einleitung zum Bericht mit einigen Worten davor."
CAPTION = "Abbildung 1: Umsatz je Quartal"
OUTRO = "Schluss des Berichts mit weiteren Worten danach."


# --- Hilfen: Bilder, PDF, DOCX ------------------------------------------------------


def picture(width=400, height=300, seed=1, mode="RGB") -> Image.Image:
    """Diagrammartiges Bild; ``seed`` macht es unterscheidbar (anderer Hash)."""
    image = Image.new(mode, (width, height), "white")
    draw = ImageDraw.Draw(image)
    for i in range(4):
        top = (seed * 37 + i * 53) % max(1, height // 2)
        draw.rectangle([20 + i * 60, top, 60 + i * 60, height - 1], fill=(30 * i, 90, 160))
    draw.text((10, 5), f"Bild {seed}", fill="black")
    return image


def png_bytes(image: Image.Image) -> bytes:
    out = io.BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_pdf(pages) -> bytes:
    """Roh-PDF; ``pages`` = Liste von Seiten, je Seite eine Liste aus
    ``("text", zeile)`` und ``("image", schlüssel)``. ``images`` (Modul-Attribut
    über ``build_pdf.images``) ordnet Schlüssel -> Pillow-Bild; gleicher
    Schlüssel = dasselbe Bildobjekt auf allen Seiten."""
    images: dict = build_pdf.images
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog = add(b"")  # 1, später gesetzt
    pages_obj = add(b"")  # 2
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    image_ids = {}
    for key, image in images.items():
        rgb = image.convert("RGB")
        raw = zlib.compress(rgb.tobytes())
        head = (
            f"<< /Type /XObject /Subtype /Image /Width {rgb.width} /Height {rgb.height} "
            f"/ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /FlateDecode "
            f"/Length {len(raw)} >>\nstream\n"
        ).encode()
        image_ids[key] = add(head + raw + b"\nendstream")
    kids = []
    for items in pages:
        ops, y, used = [], 760, []
        for kind, value in items:
            if kind == "text":
                ops.append(f"BT /F1 11 Tf 72 {y} Td ({_escape(value)}) Tj ET")
                y -= 16
            else:
                used.append(value)
                ops.append(f"q 200 0 0 150 72 {y - 150} cm /Im{value} Do Q")
                y -= 160
        content = "\n".join(ops).encode("cp1252")
        stream = add(f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream")
        xobjects = " ".join(f"/Im{k} {image_ids[k]} 0 R" for k in dict.fromkeys(used))
        page = add(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << "
            f"/Font << /F1 {font} 0 R >> /XObject << {xobjects} >> >> "
            f"/Contents {stream} 0 R >>".encode()
        )
        kids.append(page)
    objects[catalog - 1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[pages_obj - 1] = (
        f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>"
    ).encode()
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


build_pdf.images = {}


def build_docx(items) -> bytes:
    """DOCX aus ``("text", absatz)`` und ``("image", Pillow-Bild)`` in dieser Reihenfolge."""
    from docx import Document as DocxDocument
    from docx.shared import Cm

    doc = DocxDocument()
    for kind, value in items:
        if kind == "text":
            doc.add_paragraph(value)
        else:
            doc.add_picture(io.BytesIO(png_bytes(value)), width=Cm(8))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def write(tmp_path, name, data) -> str:
    path = tmp_path / name
    path.write_bytes(data)
    return str(path)


def completion(text, finish="stop"):
    return {
        "choices": [
            {"index": 0, "message": {"role": "assistant", "content": text}, "finish_reason": finish}
        ]
    }


def numbered_answers(route_calls):
    """Antwort „Beschreibung n“ je Aufruf (n ab 1)."""

    def respond(request):
        route_calls.append(json.loads(request.content))
        return httpx.Response(200, json=completion(f"Beschreibung {len(route_calls)}."))

    return respond


def sent_image(body) -> tuple[str, bytes]:
    url = body["messages"][0]["content"][1]["image_url"]["url"]
    head, _, data = url.partition(",")
    return head, base64.b64decode(data)


def fake_vectors(texts):
    return [[1.0] + [0.0] * (EMBEDDING_DIMENSIONS - 1) for _ in texts]


# --- Fixtures ---------------------------------------------------------------------------


@pytest.fixture
def media(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path / "media"
    return settings.MEDIA_ROOT


@pytest.fixture
def lmstudio():
    return Provider.objects.create(
        name="LM Studio",
        kind=Provider.Kind.OPENAI_COMPAT,
        base_url=LMSTUDIO,
        is_local=True,
        reported_models=[VISION, "allenai/olmocr-2-7b", "openai/gpt-oss-20b"],
    )


@pytest.fixture
def vision(lmstudio):
    model = AIModel.objects.create(
        provider=lmstudio, model_id=VISION, display_name="Qwen3-VL", active=False
    )
    cfg = RagSettings.load()
    cfg.describe_figures = True
    cfg.figure_model = model
    cfg.save()
    return model


@pytest.fixture
def embed():
    with mock.patch.object(ingest, "_embed", side_effect=fake_vectors) as m:
        yield m


@pytest.fixture
def owner(db):
    return User.objects.create_user("anna", password=PASSWORD, role=Role.objects.get(key="adult"))


@pytest.fixture
def collection(owner):
    return Collection.objects.create(owner=owner, name="Berichte")


def make_document(collection, name, data) -> Document:
    doc = Document(collection=collection, title=name)
    doc.file.save(name, io.BytesIO(data), save=False)
    doc.save()
    return doc


def collector(**kwargs) -> figures.Collector:
    return figures.Collector(**kwargs)


def report_pdf() -> bytes:
    build_pdf.images = {"A": picture(seed=1), "B": picture(seed=2)}
    return build_pdf(
        [
            [
                ("text", INTRO),
                ("image", "A"),
                ("text", CAPTION),
                ("text", "Text zwischen den Bildern mit genug Buchstaben."),
                ("image", "B"),
                ("text", OUTRO),
            ]
        ]
    )


# --- Aussortieren -------------------------------------------------------------------------


def test_filter_small_line_and_duplicate():
    c = collector(min_edge=150)
    assert c.add(picture(100, 100), 1) is None  # zu klein
    assert c.add(picture(160, 160), 1) is None  # Kante reicht, Fläche nicht
    assert c.add(picture(1200, 150), 1) is None  # Linie/Balken (Seitenverhältnis 8)
    assert c.add(picture(400, 300, seed=3), 1) is not None
    assert c.add(picture(400, 300, seed=3), 2) is None  # gleiches Bild (Hash)
    assert c.add(picture(400, 300, seed=4), 2) is not None
    assert len(c.figures) == 2 and c.skipped == 4


def test_prepare_scales_down_and_picks_format():
    big = picture(3000, 2000)
    data, mime = figures.prepare(big, 1024)
    out = Image.open(io.BytesIO(data))
    assert max(out.size) == 1024 and mime == "image/png"  # wenige Farben -> PNG
    noisy = Image.merge("RGB", [Image.effect_noise((800, 600), 64) for _ in range(3)])
    data, mime = figures.prepare(noisy, 1024)
    assert mime == "image/jpeg" and Image.open(io.BytesIO(data)).size == (800, 600)


def test_exif_removed_and_orientation_applied():
    image = picture(400, 300, mode="RGB")
    exif = Image.Exif()
    exif[0x010F] = "Testkamera GmbH"  # Make
    exif[0x0112] = 6  # Orientation: 90° gedreht
    exif[0x8825] = {2: (52.0, 31.0, 0.0)}  # GPS-Breite
    raw = io.BytesIO()
    image.save(raw, "JPEG", exif=exif, quality=95)
    loaded = figures.open_image(raw.getvalue())
    assert loaded.getexif()[0x010F] == "Testkamera GmbH"
    data, _mime = figures.prepare(loaded, 1024)
    out = Image.open(io.BytesIO(data))
    assert not out.getexif() and b"Testkamera" not in data and "exif" not in out.info
    assert out.size == (300, 400)  # Ausrichtung übernommen


def test_decompression_bomb_is_skipped(monkeypatch):
    monkeypatch.setattr(figures, "MAX_PIXELS", 10_000)
    assert figures.open_image(png_bytes(picture(400, 300))) is None
    c = collector()
    assert c.add_bytes(png_bytes(picture(400, 300)), None) is None


# --- Platzhalter, Position, Bildunterschrift ---------------------------------------------


def test_insert_markers_as_own_paragraphs():
    text = "Zeile eins\nZeile zwei\n\nAbsatz zwei"
    out = figures.insert_markers(text, [(4, "M1"), (None, "M2"), (0, "M0")])
    assert out == "M0\n\nZeile eins\n\nM1\n\nZeile zwei\n\nAbsatz zwei\n\nM2"
    assert figures.replace_markers("A\n\n0\n\nB", {}) == "A\n\nB"
    assert figures.replace_markers("A\n\n0\n\nB", {0: "[Abbildung: x]"}) == (
        "A\n\n[Abbildung: x]\n\nB"
    )


def test_clean_description():
    text = figures.clean_description("**Diagramm**\n\n- zeigt [Werte]\n1 ")
    assert text == "Diagramm - zeigt (Werte)"
    assert len(figures.clean_description("Wort " * 1000)) <= figures.MAX_DESCRIPTION_CHARS


def test_language_detection():
    assert figures._language("The report shows that the revenue of the company is " * 3) == (
        "Englisch"
    )
    assert figures._language("Der Bericht zeigt, dass die Kosten und die Erträge " * 3) == (
        "Deutsch"
    )
    assert figures._language("123 456") == "Deutsch"


@respx.mock
def test_pdf_order_position_and_caption(tmp_path, vision):
    calls: list = []
    respx.post(f"{LMSTUDIO}/chat/completions").mock(side_effect=numbered_answers(calls))
    path = write(tmp_path, "bericht.pdf", report_pdf())
    c = figures.collector_for()
    pages = extract.extract(path, extract.KIND_PDF, figures=c, ocr=lambda *_: "")
    assert len(c.figures) == 2 and "" in pages[0].text
    assert c.describe(pages) == 2

    paragraphs = chunking._paragraphs(pages[0].text)
    figure_paragraphs = [p for p in paragraphs if p.startswith("[Abbildung:")]
    assert figure_paragraphs == [
        f"[Abbildung: {CAPTION} – Beschreibung 1.]",
        "[Abbildung: Beschreibung 2.]",
    ]
    # Reihenfolge im Text: Einleitung, Bild A, Unterschrift, Zwischentext, Bild B, Schluss.
    text = pages[0].text
    order = [
        text.index(INTRO),
        text.index("Beschreibung 1."),
        text.index(f"\n{CAPTION}"),
        text.index("Text zwischen"),
        text.index("Beschreibung 2."),
        text.index(OUTRO),
    ]
    assert order == sorted(order)
    assert "" not in text
    # Prompt: deutsch, niedrige Temperatur, Bildunterschrift nur beim ersten Bild.
    assert calls[0]["temperature"] == figures.DESCRIBE_TEMPERATURE
    assert calls[0]["model"] == VISION
    prompt = calls[0]["messages"][0]["content"][0]["text"]
    assert "2 bis 6 Sätzen" in prompt and "wörtlich" in prompt and "Deutsch" in prompt
    assert CAPTION in prompt
    assert CAPTION not in calls[1]["messages"][0]["content"][0]["text"]
    head, data = sent_image(calls[0])
    assert head == "data:image/png;base64" and Image.open(io.BytesIO(data)).size == (400, 300)


def test_pdf_logo_on_every_page_described_once(tmp_path):
    build_pdf.images = {"L": picture(300, 300, seed=9)}
    pdf = build_pdf([[("text", INTRO), ("image", "L")] for _ in range(3)])
    c = collector()
    pages = extract.extract(write(tmp_path, "logo.pdf", pdf), extract.KIND_PDF, figures=c)
    assert len(c.figures) == 1
    assert ["" in p.text for p in pages] == [True, False, False]


def test_pdf_limits_per_page_and_document(tmp_path):
    build_pdf.images = {str(i): picture(seed=i) for i in range(6)}
    pdf = build_pdf(
        [
            [("text", INTRO), ("image", "0"), ("image", "1"), ("image", "2")],
            [("text", INTRO), ("image", "3"), ("image", "4"), ("image", "5")],
        ]
    )
    path = write(tmp_path, "viele.pdf", pdf)
    c = collector(max_per_page=2, max_per_document=3)
    extract.extract(path, extract.KIND_PDF, figures=c)
    assert [f.page for f in c.figures] == [1, 1, 2]


def test_ocr_pages_get_no_figures(tmp_path):
    """Seiten ohne Textebene: Text per OCR, das Bild ist die Seite selbst."""
    build_pdf.images = {"S": picture(seed=5)}
    path = write(tmp_path, "scan.pdf", build_pdf([[("image", "S")]]))
    c = collector()
    pages = extract.extract(
        path, extract.KIND_PDF, figures=c, ocr=lambda *_: "Erkannter Text der Seite eins."
    )
    assert pages[0].ocr and not c.figures


@respx.mock
def test_docx_order_and_caption(tmp_path, vision):
    calls: list = []
    respx.post(f"{LMSTUDIO}/chat/completions").mock(side_effect=numbered_answers(calls))
    data = build_docx(
        [
            ("text", INTRO),
            ("image", picture(seed=1)),
            ("text", CAPTION),
            ("text", "Mittlerer Absatz."),
            ("image", picture(seed=2)),
            ("image", picture(seed=2)),  # doppelt
            ("image", picture(60, 60, seed=3)),  # Symbol
            ("text", OUTRO),
        ]
    )
    c = figures.collector_for()
    pages = extract.extract(write(tmp_path, "bericht.docx", data), extract.KIND_DOCX, figures=c)
    assert c.describe(pages) == 2 and len(calls) == 2
    assert chunking._paragraphs(pages[0].text) == [
        INTRO,
        f"[Abbildung: {CAPTION} – Beschreibung 1.]",
        CAPTION,
        "Mittlerer Absatz.",
        "[Abbildung: Beschreibung 2.]",
        OUTRO,
    ]


def test_docx_limit_per_document(tmp_path):
    data = build_docx([("image", picture(seed=i)) for i in range(5)])
    c = collector(max_per_document=3)
    extract.extract(write(tmp_path, "bilder.docx", data), extract.KIND_DOCX, figures=c)
    assert len(c.figures) == 3


# --- Beschreiben: Fehler, Abbruch, Einstellungen ----------------------------------------


@respx.mock
def test_rejected_figure_is_dropped(tmp_path, vision, caplog):
    respx.post(f"{LMSTUDIO}/chat/completions").mock(
        side_effect=[
            httpx.Response(400, json={"error": {"message": "image not supported"}}),
            httpx.Response(200, json=completion("Ein Balkendiagramm.")),
        ]
    )
    c = figures.collector_for(document_id=4711)
    pages = extract.extract(write(tmp_path, "b.pdf", report_pdf()), extract.KIND_PDF, figures=c)
    assert c.describe(pages) == 1
    text = pages[0].text
    assert text.count("[Abbildung:") == 1 and "Ein Balkendiagramm." in text
    assert "" not in text and "\n\n\n" not in text
    assert "4711" in caplog.text and "HTTP 400" in caplog.text
    assert "Umsatz" not in caplog.text  # keine Inhalte im Log


@respx.mock
def test_truncated_answer_drops_figure(tmp_path, vision):
    respx.post(f"{LMSTUDIO}/chat/completions").mock(
        return_value=httpx.Response(200, json=completion("Abgeschnit", finish="length"))
    )
    c = figures.collector_for()
    pages = extract.extract(write(tmp_path, "b.pdf", report_pdf()), extract.KIND_PDF, figures=c)
    assert c.describe(pages) == 0 and "[Abbildung" not in pages[0].text


@respx.mock
def test_unreachable_provider_lets_job_wait(collection, media, vision, embed):
    respx.post(f"{LMSTUDIO}/chat/completions").mock(side_effect=httpx.ConnectError("refused"))
    doc = make_document(collection, "bericht.pdf", report_pdf())
    jobs.enqueue_index(doc)
    job = jobs.claim_next()
    assert jobs.run_job(job) == Job.Status.PENDING
    job.refresh_from_db()
    doc.refresh_from_db()
    assert job.attempts == 0 and job.run_after > job.created
    assert doc.status == Document.Status.PENDING and doc.error_text.startswith("Wartet:")
    assert not embed.called


@respx.mock
def test_server_error_is_retryable(tmp_path, vision):
    respx.post(f"{LMSTUDIO}/chat/completions").mock(return_value=httpx.Response(503))
    c = figures.collector_for()
    pages = extract.extract(write(tmp_path, "b.pdf", report_pdf()), extract.KIND_PDF, figures=c)
    with pytest.raises(figures.FigureError) as exc:
        c.describe(pages)
    assert exc.value.retryable and not exc.value.unreachable


@respx.mock
def test_cancel_between_figures(collection, media, vision, embed):
    calls: list = []
    respx.post(f"{LMSTUDIO}/chat/completions").mock(side_effect=numbered_answers(calls))
    doc = make_document(collection, "bericht.pdf", report_pdf())
    with pytest.raises(extract.Interrupted):
        ingest.index_document(doc, should_stop=lambda: bool(calls))
    assert len(calls) == 1 and not embed.called


@respx.mock
def test_index_document_counts_figures(collection, media, vision, embed):
    calls: list = []
    respx.post(f"{LMSTUDIO}/chat/completions").mock(side_effect=numbered_answers(calls))
    doc = make_document(collection, "bericht.pdf", report_pdf())
    result = ingest.index_document(doc)
    doc.refresh_from_db()
    assert result.figures == 2 and doc.figures_described == 2
    texts = " ".join(doc.chunks.values_list("text", flat=True))
    assert "[Abbildung: Abbildung 1: Umsatz je Quartal – Beschreibung 1.]" in texts


@respx.mock
def test_settings_off_makes_no_calls(collection, media, lmstudio, embed):
    route = respx.post(f"{LMSTUDIO}/chat/completions")
    doc = make_document(collection, "bericht.pdf", report_pdf())
    assert figures.collector_for() is None
    result = ingest.index_document(doc)
    assert result.figures == 0 and route.call_count == 0
    texts = " ".join(doc.chunks.values_list("text", flat=True))
    assert "Abbildung:" not in texts and "" not in texts


def test_missing_model_is_permanent_error(tmp_path):
    cfg = RagSettings.load()
    cfg.describe_figures = True
    cfg.save()
    c = figures.collector_for()
    pages = extract.extract(write(tmp_path, "b.pdf", report_pdf()), extract.KIND_PDF, figures=c)
    with pytest.raises(figures.FigureError) as exc:
        c.describe(pages)
    assert not exc.value.retryable and "kein Modell" in exc.value.message


# --- Bilddateien als Dokument ---------------------------------------------------------------


def _encoded(fmt, **kwargs) -> bytes:
    out = io.BytesIO()
    picture().save(out, fmt, **kwargs)
    return out.getvalue()


@pytest.mark.parametrize(
    ("fmt", "name"),
    [("JPEG", "foto.jpg"), ("JPEG", "foto.JPEG"), ("PNG", "bild.png"), ("TIFF", "scan.tif")]
    + [("TIFF", "scan.tiff"), ("WEBP", "bild.webp")],
)
def test_detect_kind_images_by_magic_bytes(fmt, name):
    assert extract.detect_kind(io.BytesIO(_encoded(fmt)), name) == extract.KIND_IMAGE


@pytest.mark.parametrize(
    ("data", "name", "fragment"),
    [
        (_encoded("PNG"), "foto.jpg", "erkannt: PNG-Bild"),
        (_encoded("JPEG"), "bild.png", "erkannt: JPEG-Bild"),
        (b"Nur Text, kein Bild", "bild.png", "passt nicht zur Dateiendung"),
        (_encoded("GIF"), "anim.gif", "nicht unterstützt"),
        (b"\x00\x00\x00\x18ftypheic\x00\x00\x00\x00", "foto.heic", "nicht unterstützt"),
    ],
)
def test_detect_kind_rejects_wrong_images(data, name, fragment):
    with pytest.raises(extract.UnsupportedFile) as exc:
        extract.detect_kind(io.BytesIO(data), name)
    assert fragment in str(exc.value)


def test_upload_image_document(client, owner, collection, media):
    client.force_login(owner)
    url = reverse("chat:api_collection_documents", args=[collection.pk])
    ok = client.post(url, {"file": SimpleUploadedFile("foto.jpg", _encoded("JPEG"))})
    assert ok.status_code == 201, ok.content
    doc = Document.objects.get(pk=ok.json()["id"])
    assert doc.file.name.endswith(".jpg") and Job.objects.count() == 1
    wrong = client.post(url, {"file": SimpleUploadedFile("foto.jpg", _encoded("PNG"))})
    assert wrong.status_code == 415 and "PNG-Bild" in wrong.json()["error"]
    fake = client.post(url, {"file": SimpleUploadedFile("foto.png", b"<html>kein Bild</html>")})
    assert fake.status_code == 415
    assert Document.objects.count() == 1


def test_multipage_tiff_reads_each_page(tmp_path):
    out = io.BytesIO()
    first, second = picture(seed=1), picture(seed=2)
    first.save(out, "TIFF", save_all=True, append_images=[second])
    seen = []

    def fake_ocr(pdf_path, page):
        seen.append((Path(pdf_path).read_bytes()[:5], page))
        return f"Seite {len(seen)} Text"

    c = collector()
    pages = extract.extract(
        write(tmp_path, "scan.tiff", out.getvalue()), extract.KIND_IMAGE, ocr=fake_ocr, figures=c
    )
    assert [p.number for p in pages] == [1, 2] and all(p.ocr for p in pages)
    assert seen == [(b"%PDF-", 1), (b"%PDF-", 1)]
    assert pages[0].text == "0\n\nSeite 1 Text" and len(c.figures) == 2


@respx.mock
def test_image_document_ocr_and_description_without_exif(collection, media, vision, embed):
    calls: list = []
    respx.post(f"{LMSTUDIO}/chat/completions").mock(side_effect=numbered_answers(calls))
    exif = Image.Exif()
    exif[0x010F] = "Testkamera GmbH"
    raw = io.BytesIO()
    picture(1600, 1200).save(raw, "JPEG", exif=exif, quality=90)
    doc = make_document(collection, "foto.jpg", raw.getvalue())
    result = ingest.index_document(doc, ocr=lambda *_: "Schild am Eingang: Rathaus")
    assert result.figures == 1 and result.ocr_pages == 1 and result.pages == 1
    text = " ".join(doc.chunks.values_list("text", flat=True))
    assert text == "[Abbildung: Beschreibung 1.]\n\nSchild am Eingang: Rathaus"
    assert doc.chunks.get().page is None
    head, data = sent_image(calls[0])
    sent = Image.open(io.BytesIO(data))
    assert max(sent.size) == 1024 and not sent.getexif() and b"Testkamera" not in data


def test_image_without_ocr_programs(tmp_path):
    def missing(*_):
        raise extract.OcrUnavailable("tesseract")

    with pytest.raises(extract.ExtractionError) as exc:
        extract.extract(write(tmp_path, "a.png", _encoded("PNG")), extract.KIND_IMAGE, ocr=missing)
    assert "nicht installiert" in str(exc.value)


def test_broken_image_file(tmp_path):
    path = write(tmp_path, "kaputt.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 50)
    with pytest.raises(extract.ExtractionError):
        extract.extract(path, extract.KIND_IMAGE, ocr=lambda *_: "")


# --- Einstellungen und Admin ------------------------------------------------------------------


def test_settings_validation(lmstudio):
    from django.core.exceptions import ValidationError

    cfg = RagSettings.load()
    cfg.describe_figures = True
    with pytest.raises(ValidationError) as exc:
        cfg.clean()
    assert "figure_model" in exc.value.message_dict
    anthropic = Provider.objects.create(name="Anthropic", kind=Provider.Kind.ANTHROPIC)
    cfg.figure_model = AIModel.objects.create(
        provider=anthropic, model_id="claude-x", display_name="Claude", supports_vision=True
    )
    with pytest.raises(ValidationError):
        cfg.clean()
    cfg.figure_model = AIModel.objects.create(provider=lmstudio, model_id=VISION, display_name="Q")
    cfg.clean()


def test_figure_model_choices(lmstudio):
    cloud = Provider.objects.create(
        name="OpenAI", kind=Provider.Kind.OPENAI_COMPAT, reported_models=["gpt-4o", "gpt-5"]
    )
    AIModel.objects.create(
        provider=cloud, model_id="gpt-4o", display_name="GPT-4o", supports_vision=True
    )
    AIModel.objects.create(provider=cloud, model_id="o1-mini", display_name="o1 mini")
    choices = settings_form.model_choices(settings_form.FIGURE)
    groups = dict(choices[1:])
    recommended = [label for _value, label in groups[settings_form.RECOMMENDED]]
    assert any(VISION in label for label in recommended)
    assert any("GPT-4o" in label for label in recommended)
    assert not any("olmocr" in label for label in recommended)
    flat = [label for _g, options in choices[1:] for _v, label in options]
    assert not any("o1 mini" in label for label in flat)  # Cloud ohne Bild-Eingabe
    # olmOCR ist Hauptart OCR (reine Texterkennung) und für Abbildungen nicht wählbar.
    assert not any("olmocr" in label for label in flat)


def test_new_figure_model_is_created_vision_only(lmstudio):
    model = settings_form.resolve_choice(settings_form.FIGURE, f"r{lmstudio.pk}:{VISION}")
    saved = settings_form.materialize(model)
    assert saved.pk and saved.supports_vision and not saved.active


@pytest.fixture
def admin_client(client):
    user = User.objects.create_superuser(
        username="jo", password=PASSWORD, email="jo@example.invalid"
    )
    client.force_login(user)
    return client


def test_overview_shows_figure_settings(admin_client, vision):
    response = admin_client.get(reverse("admin:rag_ragoverview_changelist"))
    body = response.content.decode()
    assert response.status_code == 200
    assert "Abbildungen" in body and "beschreiben mit" in body and VISION in body


def test_overview_figures_off(admin_client):
    response = admin_client.get(reverse("admin:rag_ragoverview_changelist"))
    assert "nicht beschreiben" in response.content.decode()


def test_settings_page_has_figure_fields(admin_client, vision):
    url = reverse("admin:rag_ragsettingsproxy_change", args=[RagSettings.SINGLETON_PK])
    body = admin_client.get(url).content.decode()
    for name in ("describe_figures", "figure_model", "figure_max_per_document", "figure_min_edge"):
        assert f'name="{name}"' in body
    assert "gehen sie an den Anbieter" in body
