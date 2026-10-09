"""Erzeugt die kleinen Beispieldateien für tests/test_ingest.py.

Aufruf (einmalig, Ergebnis liegt im Repository):
    .venv/bin/python tests/data/make_samples.py

- sample.pdf: Seite 1 und 3 mit Textebene, Seite 2 leer (wie eine gescannte
  Seite ohne Text -> OCR-Pfad).
- scanned.pdf: eine Seite, die nur aus einem Bild von Text besteht (für den
  OCR-Integrationstest; gerendert mit pdftoppm aus einer Textseite).
- sample.docx: Überschrift, Absätze, Tabelle (python-docx).
- sample.txt (UTF-8), sample_cp1252.txt, sample.md.
"""

import subprocess
import tempfile
import zlib
from pathlib import Path

HERE = Path(__file__).resolve().parent


def _pdf(objects: list[bytes]) -> bytes:
    """Minimales PDF aus nummerierten Objekten (1 = Catalog) mit xref-Tabelle."""
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def _stream(data: bytes, extra: str = "") -> bytes:
    return f"<< /Length {len(data)} {extra}>>\nstream\n".encode() + data + b"\nendstream"


def _text_content(lines: list[str]) -> bytes:
    ops = ["BT", "/F1 14 Tf", "72 760 Td", "18 TL"]
    for line in lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops.append(f"({safe}) '")
    ops.append("ET")
    # WinAnsiEncoding: Umlaute als cp1252-Bytes.
    return "\n".join(ops).encode("cp1252")


def make_text_pdf(pages: list[list[str]]) -> bytes:
    n = len(pages)
    # 1 Catalog, 2 Pages, 3 Font, dann je Seite: Page, Content
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    for i, lines in enumerate(pages):
        content_id = 5 + 2 * i
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        objects.append(_stream(_text_content(lines)))
    return _pdf(objects)


def make_image_pdf(gray: bytes, width: int, height: int) -> bytes:
    data = zlib.compress(gray)
    content = b"q 612 0 0 792 0 0 cm /Im1 Do Q"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /XObject << /Im1 4 0 R >> >> /Contents 5 0 R >>",
        _stream(
            data,
            f"/Type /XObject /Subtype /Image /Width {width} /Height {height} "
            f"/ColorSpace /DeviceGray /BitsPerComponent 8 /Filter /FlateDecode ",
        ),
        _stream(content),
    ]
    return _pdf(objects)


def _read_pgm(path: Path) -> tuple[bytes, int, int]:
    raw = path.read_bytes()
    # P5\n<w> <h>\n255\n<daten>
    parts = raw.split(b"\n", 3)
    assert parts[0] == b"P5"
    width, height = map(int, parts[1].split())
    return parts[3], width, height


PAGE1 = [
    "Hausordnung der Familie Beispiel",
    "Die Waschmaschine darf nur zwischen 8 und 20 Uhr laufen.",
    "Der Müll wird jeden Dienstag an die Straße gestellt.",
]
PAGE3 = [
    "Anhang: Notfallnummern",
    "Feuerwehr und Rettungsdienst erreichen Sie unter 112.",
    "Der Hausmeister heißt Herr Größer und wohnt im Erdgeschoss.",
]
SCAN = [
    "Mietvertrag Seite 1",
    "Die Kaltmiete beträgt 850 Euro im Monat.",
    "Die Kaution beträgt drei Monatsmieten.",
]


def main() -> None:
    (HERE / "sample.pdf").write_bytes(make_text_pdf([PAGE1, [], PAGE3]))

    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "scan-src.pdf"
        src.write_bytes(make_text_pdf([[line for line in SCAN]]))
        subprocess.run(
            ["pdftoppm", "-r", "150", "-gray", "-singlefile", str(src), str(Path(tmp) / "p")],
            check=True,
        )
        gray, w, h = _read_pgm(Path(tmp) / "p.pgm")
        (HERE / "scanned.pdf").write_bytes(make_image_pdf(gray, w, h))

    from docx import Document

    doc = Document()
    doc.add_heading("Rezeptsammlung", level=1)
    doc.add_paragraph("Apfelkuchen nach Omas Art braucht säuerliche Äpfel.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Zutat"
    table.cell(0, 1).text = "Menge"
    table.cell(1, 0).text = "Mehl"
    table.cell(1, 1).text = "500 g"
    doc.add_paragraph("Bei 180 Grad 45 Minuten backen.")
    doc.save(HERE / "sample.docx")

    text = "Einkaufsliste für Samstag\n\nMilch, Brot, Käse und Äpfel.\n"
    (HERE / "sample.txt").write_text(text, encoding="utf-8")
    (HERE / "sample_cp1252.txt").write_bytes(text.encode("cp1252"))
    (HERE / "sample.md").write_text(
        "# Urlaubsplanung\n\n- Zelt einpacken\n- **Sonnencreme** nicht vergessen\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
